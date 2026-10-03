from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pypdf import PdfReader

SUPPORTED_SUFFIXES = {
    ".pdf",
    ".xlsx",
    ".xls",
    ".pptx",
    ".ppt",
    ".docx",
    ".doc",
    ".csv",
    ".txt",
    ".md",
    ".rtf",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
}

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".log", ".rtf"}
OLE_SUFFIXES = {".ppt", ".doc", ".xls"}


def parse_file(path: Path, max_pages: int | None = None) -> dict[str, Any]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path, max_pages=max_pages)
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix in {".txt", ".md", ".log", ".json", ".xml", ".html", ".htm", ".rtf"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        pages = _page_texts(text)
        return {"text": text, "pages": max(1, len(pages)), "kind": "text", "tables": [], "page_texts": pages, "ocr_used": False}
    if suffix in OLE_SUFFIXES:
        return _parse_ole(path, suffix)
    if _looks_text(path):
        text = path.read_text(encoding="utf-8", errors="ignore")
        pages = _page_texts(text)
        return {"text": text, "pages": max(1, len(pages)), "kind": "text", "tables": [], "page_texts": pages, "ocr_used": False}
    return {"text": "", "pages": 0, "kind": "unsupported", "tables": [], "page_texts": [], "ocr_used": False, "error": f"unsupported type {suffix}"}


def _parse_pdf(path: Path, max_pages: int | None = None) -> dict[str, Any]:
    reader = PdfReader(str(path))
    total = len(reader.pages)
    limit = total if not max_pages else min(total, max(1, max_pages))
    pages = [(reader.pages[index].extract_text() or "") for index in range(limit)]
    ocr_used = False
    if _needs_ocr(pages):
        ocr_pages = _ocr_pdf(path, max_pages=limit)
        if ocr_pages:
            for index, text in enumerate(ocr_pages):
                if index >= len(pages):
                    pages.append(text)
                elif len((pages[index] or "").strip()) < 40:
                    pages[index] = text
            ocr_used = True
    return {
        "text": "\n\n".join(pages),
        "pages": total,
        "kind": "pdf",
        "tables": [],
        "page_texts": pages,
        "ocr_used": ocr_used,
    }


def _parse_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[dict[str, Any]] = []
    parts: list[str] = []
    page_texts: list[str] = []
    sheet_count = 0
    for sheet in workbook.worksheets:
        sheet_count += 1
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_cell(value) for value in rows[0]]
        body = []
        for raw in rows[1:201]:
            record = {headers[index] or f"col_{index}": _cell(raw[index] if index < len(raw) else "") for index in range(len(headers))}
            if any(str(value).strip() for value in record.values()):
                body.append(record)
        tables.append({"sheet": sheet.title, "headers": headers, "rows": body})
        block = [f"# Sheet {sheet.title}", " | ".join(headers)]
        for record in body[:40]:
            block.append(" | ".join(str(record.get(header, "")) for header in headers))
        page_texts.append("\n".join(block))
        parts.extend(block)
    workbook.close()
    return {
        "text": "\n".join(parts),
        "pages": sheet_count or len(tables),
        "kind": "spreadsheet",
        "tables": tables,
        "page_texts": page_texts,
        "ocr_used": False,
    }


def _parse_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    deck = Presentation(str(path))
    slides: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        bits = [f"# Slide {index}"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                bits.append(shape.text)
        slides.append("\n".join(bits))
    return {"text": "\n\n".join(slides), "pages": len(slides), "kind": "presentation", "tables": [], "page_texts": slides, "ocr_used": False}


def _parse_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(path))
    paragraphs = [item.text for item in document.paragraphs if item.text]
    text = "\n".join(paragraphs)
    pages = _page_texts(text, size=1600)
    return {"text": text, "pages": max(1, len(pages)), "kind": "document", "tables": [], "page_texts": pages, "ocr_used": False}


def _parse_csv(path: Path) -> dict[str, Any]:
    import csv

    with path.open(newline="", encoding="utf-8", errors="ignore") as handle:
        reader = csv.DictReader(handle)
        rows = [dict(row) for _, row in zip(range(200), reader)]
        headers = reader.fieldnames or []
    lines = [" | ".join(headers)]
    for row in rows[:40]:
        lines.append(" | ".join(str(row.get(header, "")) for header in headers))
    text = "\n".join(lines)
    return {
        "text": text,
        "pages": 1,
        "kind": "spreadsheet",
        "tables": [{"sheet": path.stem, "headers": headers, "rows": rows}],
        "page_texts": [text],
        "ocr_used": False,
    }


def _parse_ole(path: Path, suffix: str) -> dict[str, Any]:
    if suffix == ".xls":
        try:
            return _parse_xlsx(path)
        except Exception:
            pass
    text = _ole_strings(path)
    kind = "presentation" if suffix == ".ppt" else "document" if suffix == ".doc" else "spreadsheet"
    pages = _page_texts(text)
    return {
        "text": text,
        "pages": max(1, len(pages) or (1 if text else 0)),
        "kind": kind,
        "tables": [],
        "page_texts": pages or ([text] if text else []),
        "ocr_used": False,
    }


def _ole_strings(path: Path) -> str:
    data = path.read_bytes()
    ascii_chunks = [chunk.decode("ascii", "ignore") for chunk in re.findall(rb"[\x20-\x7e]{8,}", data)]
    utf16_chunks = []
    try:
        decoded = data.decode("utf-16le", errors="ignore")
        utf16_chunks = re.findall(r"[\w][\w \t.,:;/-]{7,}", decoded)
    except Exception:
        pass
    parts = ascii_chunks[:120] + utf16_chunks[:80]
    return "\n".join(part.strip() for part in parts if part.strip())


def _looks_text(path: Path) -> bool:
    try:
        sample = path.read_bytes()[:8192]
    except OSError:
        return False
    if not sample or b"\x00" in sample:
        return False
    return True


def _needs_ocr(pages: list[str]) -> bool:
    if not pages:
        return True
    text = " ".join(pages)
    return len(text.strip()) < 60 * max(len(pages), 1)


def _ocr_pdf(path: Path, max_pages: int = 2) -> list[str]:
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError:
        return []
    try:
        document = pdfium.PdfDocument(str(path))
    except Exception:
        return []
    pages: list[str] = []
    limit = min(len(document), max_pages)
    for index in range(limit):
        try:
            page = document[index]
            bitmap = page.render(scale=2).to_pil()
            pages.append(pytesseract.image_to_string(bitmap) or "")
        except Exception:
            pages.append("")
    return pages


def _page_texts(text: str, size: int = 1800) -> list[str]:
    if not text:
        return []
    if "\f" in text:
        return [part.strip() for part in text.split("\f") if part.strip()]
    return [text[index : index + size] for index in range(0, len(text), size)] or [text]


def _cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
